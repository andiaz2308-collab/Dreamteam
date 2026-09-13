from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.schemas.candidate import CandidateProfile
from app.schemas.job import JobProfile
from app.schemas.optimization import CustomizedCV


def safe_filename(value: str, fallback: str = "CV_adaptado") -> str:
    cleaned = re.sub(r"[^\w\-áéíóúÁÉÍÓÚñÑ ]+", "", value or "", flags=re.UNICODE)
    cleaned = "_".join(cleaned.split())
    return cleaned[:80] or fallback


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def _add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.size = Pt(11)


def build_customized_cv_docx(
    profile: CandidateProfile,
    customized: CustomizedCV,
    job: JobProfile,
) -> bytes:
    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(profile.name)
    run.bold = True
    run.font.size = Pt(18)

    headline = customized.headline or profile.headline
    if headline:
        line = document.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = line.add_run(headline)
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x3F, 0x3F, 0x46)

    contact = [
        str(item)
        for item in [profile.location, profile.email, profile.phone]
        if item
    ]
    if contact:
        line = document.add_paragraph(" · ".join(contact))
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    target = document.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = target.add_run(f"Postulación: {job.title} — {job.company}")
    run.italic = True
    run.font.size = Pt(10)

    if customized.summary:
        _add_heading(document, "Perfil profesional")
        _add_body(document, customized.summary)

    if customized.emphasize:
        _add_heading(document, "Competencias alineadas a la oferta")
        _add_body(document, ", ".join(customized.emphasize))

    if customized.experience:
        _add_heading(document, "Experiencia")
        for item in customized.experience:
            _add_body(document, f"• {item}")

    if profile.education:
        _add_heading(document, "Educación")
        for item in profile.education:
            line = " · ".join(
                part
                for part in [item.degree, item.field, item.institution]
                if part
            )
            if line:
                _add_body(document, f"• {line}")

    if customized.skills:
        _add_heading(document, "Habilidades")
        _add_body(document, ", ".join(customized.skills))

    if profile.languages:
        _add_heading(document, "Idiomas")
        for item in profile.languages:
            line = " · ".join(part for part in [item.language, item.level] if part)
            if line:
                _add_body(document, f"• {line}")

    if profile.certifications:
        _add_heading(document, "Certificaciones")
        for item in profile.certifications:
            line = " · ".join(
                part for part in [item.name, item.institution, item.date] if part
            )
            if line:
                _add_body(document, f"• {line}")

    if profile.projects:
        _add_heading(document, "Proyectos")
        for item in profile.projects:
            line = item.name
            if item.description:
                line = f"{line}: {item.description}"
            _add_body(document, f"• {line}")

    note = document.add_paragraph()
    run = note.add_run(
        "Nota: CV adaptado a partir del CV maestro. No se inventaron datos."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_download_name(job: JobProfile, profile: CandidateProfile | None = None) -> str:
    person = safe_filename(profile.name if profile else "", "CV")
    company = safe_filename(job.company, "Empresa")
    role = safe_filename(job.title, "Cargo")
    return f"CV_{person}_{company}_{role}.docx"
