from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.services.document.document_types import DocumentContent
from app.services.document.limits import MAX_FILE_SIZE, MAX_TEXT_LENGTH
from app.services.document.text_cleaner import clean_text


class DOCXParserError(Exception):
    pass


def _paragraph_text(document: Document) -> list[str]:
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return blocks


def extract_docx(
    file_bytes: bytes,
    filename: str,
    content_type: str,
) -> DocumentContent:
    if not file_bytes:
        raise DOCXParserError("El archivo esta vacio.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise DOCXParserError("El DOCX supera el limite de 5 MB.")

    if not file_bytes.startswith(b"PK"):
        raise DOCXParserError("El archivo no parece ser un DOCX valido.")

    try:
        document = Document(BytesIO(file_bytes))
    except PackageNotFoundError as exc:
        raise DOCXParserError("No se pudo abrir el DOCX.") from exc
    except Exception as exc:
        raise DOCXParserError("No se pudo abrir el DOCX.") from exc

    raw_text = "\n\n".join(_paragraph_text(document))
    cleaned_text = clean_text(raw_text)

    if len(cleaned_text) > MAX_TEXT_LENGTH:
        raise DOCXParserError("El texto extraido supera el limite permitido.")

    return DocumentContent(
        filename=filename,
        content_type=content_type,
        extension=".docx",
        text=cleaned_text,
        pages=None,
        characters=len(cleaned_text),
        extraction_method="python-docx",
        has_text=bool(cleaned_text),
    )
